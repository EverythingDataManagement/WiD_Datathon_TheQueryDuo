#!/usr/bin/env python3
# =====================================================================
#  build_fragility_index.py  —  Query Queens, GROW track  (FINALE)
#  Forward-looking, prescriptive resilience — all from the production
#  master, no new data, no black-box ML. Three parts:
#    1. PARTIAL-SHOCK resilience: knock out a realistic -20/-35/-50% of
#       the top producer (not just 100%) and see what's coverable.
#    2. HISTORICAL validation: the largest real 1-year drops each top
#       producer has actually had -> shocks this size have happened.
#    3. STAPLE FRAGILITY INDEX (the model): combine Exposure (top-producer
#       concentration) x Hazard (that producer's volatility) x Buffer
#       (how thin the backfill is) into a risk score per crop, ranked,
#       PLUS the countries best placed to reinforce each fragile crop.
#
#  Reads:  Master_clean.xlsx  (sheet 'master')   -- or change below
#  Writes: Fragility_Index.xlsx
#  Colab:  !pip install openpyxl --quiet
# =====================================================================

import os
import pandas as pd
import numpy as np

# ---------- SETTINGS ----------
MASTER_FILE  = "Master_clean.xlsx"   # the workbook built by 01_build_master.py
MASTER_SHEET = "master"              # its tidy tab
OUT_FILE     = "Fragility_Index.xlsx"
START, END   = 2000, 2023
SHOCK_LEVELS = [0.20, 0.35, 0.50, 1.00]   # realistic partial shocks + total
REQUIRED_COLS = ["Year", "Country", "Crop", "Production_tonnes"]


# ---------- load the master (with a clear error if anything is off) ----------
def load_master(path, sheet):
    """Load the master sheet, but fail with a message that says exactly what to fix.
    Same guard as the other scripts: most load errors are just the file not being
    in the Colab session, or a name/sheet mismatch."""
    if not os.path.exists(path):
        here = sorted(f for f in os.listdir(".") if f.lower().endswith((".xlsx", ".xls")))
        lines = [f"Could not find '{path}' in the working directory:", f"    {os.getcwd()}"]
        if here:
            lines += ["Excel files that ARE here right now:",
                      "    " + ", ".join(here),
                      f"Fix: set MASTER_FILE to one of those, or re-upload '{path}'."]
        else:
            lines += ["There are no Excel files in this directory at all.",
                      "In Colab, uploaded files disappear when the runtime restarts —",
                      "re-run 01_build_master.py (or your upload cell), then run this again."]
        raise FileNotFoundError("\n".join(lines))

    xls = pd.ExcelFile(path)
    match = next((s for s in xls.sheet_names if s.lower() == sheet.lower()), None)  # case-insensitive
    if match is None:
        raise ValueError(
            f"'{path}' opened, but it has no sheet named '{sheet}'.\n"
            f"Sheets in this file: {xls.sheet_names}\n"
            f"Fix: set MASTER_SHEET to one of those.")

    df = pd.read_excel(path, sheet_name=match)
    missing = [c for c in REQUIRED_COLS if c not in df.columns]
    if missing:
        raise KeyError(
            f"Sheet '{sheet}' is missing required column(s): {missing}\n"
            f"Columns found: {list(df.columns)}\n"
            f"Fix: point at the tidy master tab (Year / Country / Crop / Production_tonnes).")
    return df


m = load_master(MASTER_FILE, MASTER_SHEET)
m = m[m["Year"].between(START, END)]
crops = sorted(m["Crop"].unique())


# ---------- shared helpers ----------
def crop_frame(cr):
    g = m[m["Crop"] == cr]
    cur = g[g["Year"] == END].groupby("Country")["Production_tonnes"].sum()
    peak = g.groupby("Country")["Production_tonnes"].max()
    return g, cur, peak

def backfill_capacity(cur, peak, top):
    others = [c for c in cur.index if c != top]
    return (peak.reindex(others).fillna(0) - cur.reindex(others).fillna(0)).clip(lower=0)

def annual_series(country, crop):
    """Year-indexed production series for one country+crop (sums any duplicate
    rows within a year, and sorts by year so pct_change is truly chronological)."""
    return (m[(m["Country"] == country) & (m["Crop"] == crop)]
            .groupby("Year")["Production_tonnes"].sum().sort_index())

# =====================================================================
# 1. PARTIAL-SHOCK RESILIENCE
# =====================================================================
rows = []
info = {}
for cr in crops:
    g, cur, peak = crop_frame(cr)
    if cur.empty:
        continue
    top = cur.idxmax(); out = cur.max()
    bf = backfill_capacity(cur, peak, top).sum()
    info[cr] = {"top": top, "out": out, "backfill": bf}
    row = {"Crop": cr, "Top producer": top, "Top output (Mt)": round(out / 1e6, 1),
           "Backfill capacity (Mt)": round(bf / 1e6, 1)}
    for pct in SHOCK_LEVELS:
        row[f"-{int(pct*100)}% shock"] = "covered" if bf >= out * pct else "GAP"
    rows.append(row)
partial = pd.DataFrame(rows)

# =====================================================================
# 2. HISTORICAL VALIDATION — largest real 1-year drop per top producer
# =====================================================================
rows = []
for cr in crops:
    if cr not in info:
        continue
    top = info[cr]["top"]
    s = annual_series(top, cr)                      # year-indexed, chronological
    yoy = (s.pct_change() * 100).dropna()
    if yoy.empty:
        continue
    worst = yoy.min()
    worst_year = int(yoy.idxmin())                  # the YEAR of the worst drop (fixed)
    rows.append({"Crop": cr, "Top producer": top,
                 "Worst 1-yr drop %": round(worst, 0),
                 "Worst drop year": worst_year,
                 "Avg 1-yr volatility (std %)": round(yoy.std(), 1),
                 "Note": "a shock of this size has already happened"})
history = pd.DataFrame(rows)

# =====================================================================
# 3. STAPLE FRAGILITY INDEX (the model)
#    Exposure  = top-producer share of world output (%)
#    Hazard    = volatility of the top producer's output (std of YoY %)
#    Buffer    = resilience ratio (backfill / full loss); low = thin
#    Risk is HIGH when buffer is thin AND exposure/hazard are high.
# =====================================================================
def norm(x, lo, hi):
    return np.clip((x - lo) / (hi - lo + 1e-9), 0, 1)

comp = []
for cr in crops:
    if cr not in info:
        continue
    g, cur, peak = crop_frame(cr)
    top = info[cr]["top"]
    exposure = cur.max() / cur.sum() * 100
    hazard = (annual_series(top, cr).pct_change() * 100).std()
    ratio = info[cr]["backfill"] / info[cr]["out"]          # buffer
    comp.append({"Crop": cr, "Top producer": top,
                 "Exposure (top share %)": round(exposure, 0),
                 "Hazard (volatility %)": round(hazard, 1),
                 "Buffer (resilience ratio)": round(ratio, 2)})
comp = pd.DataFrame(comp)

# BUFFER IS A GATE, not just a weight. If the world can physically cover the crop
# (resilience ratio >= 1), it is LOW risk no matter how concentrated/volatile it is.
# Exposure and hazard only rank the crops that CANNOT be covered.
comp["exposure_n"] = norm(comp["Exposure (top share %)"], comp["Exposure (top share %)"].min(), comp["Exposure (top share %)"].max())
comp["hazard_n"]   = norm(comp["Hazard (volatility %)"], comp["Hazard (volatility %)"].min(), comp["Hazard (volatility %)"].max())
# buffer_n: 0 when fully coverable (ratio>=1), rising to 1 as the ratio falls to 0.
comp["buffer_n"]   = (1 - comp["Buffer (resilience ratio)"].clip(lower=0, upper=1))
# base risk is how far short the backfill falls; exposure/hazard sharpen it, but only
# matter when there IS a gap (multiplied by buffer_n so a coverable crop stays near 0).
sharpen = 0.6 * comp["exposure_n"] + 0.4 * comp["hazard_n"]
comp["Fragility score (0-100)"] = (100 * comp["buffer_n"] * (0.7 + 0.3 * sharpen)).round(0)
comp["Risk band"] = pd.cut(comp["Fragility score (0-100)"], [-1, 20, 55, 101],
                           labels=["Low", "Moderate", "High"])
index_tab = comp[["Crop", "Top producer", "Exposure (top share %)", "Hazard (volatility %)",
                  "Buffer (resilience ratio)", "Fragility score (0-100)", "Risk band"]] \
                 .sort_values("Fragility score (0-100)", ascending=False)

# validation: does the index flag the crops that already failed the partial shock?
val = partial.merge(comp[["Crop", "Fragility score (0-100)", "Risk band"]], on="Crop")
val["Fails at -35%?"] = val["-35% shock"].eq("GAP").map({True: "yes", False: "no"})
validation = val[["Crop", "Fragility score (0-100)", "Risk band", "Fails at -35%?"]] \
                 .sort_values("Fragility score (0-100)", ascending=False)

# =====================================================================
# 3b. PRESCRIPTION — who can reinforce each fragile crop?
#     The countries with the most LATENT capacity: biggest gap between
#     their historical peak and current output (excluding the top producer).
# =====================================================================
rows = []
fragile = index_tab[index_tab["Risk band"].isin(["High", "Moderate"])]["Crop"].tolist()
for cr in fragile:
    g, cur, peak = crop_frame(cr)
    top = info[cr]["top"]
    head = backfill_capacity(cur, peak, top).sort_values(ascending=False)
    for rank, (country, gap) in enumerate(head.head(3).items(), 1):
        rows.append({"Crop (fragile)": cr, "Reinforcement rank": rank,
                     "Country": country,
                     "Latent capacity (Mt)": round(gap / 1e6, 1),
                     "= peak minus current output": "countries that have produced more before"})
prescription = pd.DataFrame(rows)

# =====================================================================
# WRITE
# =====================================================================
readme = pd.DataFrame([
    ["Query Queens - STAPLE FRAGILITY INDEX (project finale)", ""],
    ["Forward-looking, prescriptive resilience - all from the production master.", ""],
    ["", ""],
    ["partial_shock", "Can the world cover a -20/-35/-50/-100% loss of each crop's #1 producer?"],
    ["historical_drops", "The largest REAL 1-year drops each top producer has had (shocks this size happen)."],
    ["fragility_index", "The model: Exposure x Hazard x Buffer -> a fragility score + risk band per crop."],
    ["index_validation", "Does the index flag the crops that already fail a realistic -35% shock? (it should)"],
    ["prescription", "For each fragile crop, the countries best placed to reinforce it (most latent capacity)."],
    ["", ""],
    ["How the index works", ""],
    ["Exposure = top producer's share of world output. Hazard = that producer's output volatility.", ""],
    ["Buffer = resilience ratio (backfill / full loss). Buffer is a GATE: a crop the world can", ""],
    ["physically cover (ratio >= 1) scores ~0, however concentrated or volatile its top producer is.", ""],
    ["For crops with a real gap:  score = 100 x buffer_shortfall x (0.7 + 0.3 x sharpen),", ""],
    ["where buffer_shortfall = 1 - min(ratio,1) and sharpen = 0.6 x Exposure + 0.4 x Hazard", ""],
    ["(exposure and hazard normalised across crops). So exposure/hazard only sharpen an existing gap.", ""],
    ["", ""],
    ["Why not a forecast model: every input is MEASURED, so the score is transparent and defensible.", ""],
    ["Validation is that it correctly flags the crops that already fail a realistic partial shock.", ""],
    ["Limitation: tonnage-based structural risk; does not model shock CAUSE or ramp-up SPEED.", ""],
], columns=["A", "B"])

with pd.ExcelWriter(OUT_FILE, engine="openpyxl") as xl:
    readme.to_excel(xl, sheet_name="read_me", index=False, header=False)
    partial.to_excel(xl, sheet_name="partial_shock", index=False)
    history.to_excel(xl, sheet_name="historical_drops", index=False)
    index_tab.to_excel(xl, sheet_name="fragility_index", index=False)
    validation.to_excel(xl, sheet_name="index_validation", index=False)
    prescription.to_excel(xl, sheet_name="prescription", index=False)

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
wb = load_workbook(OUT_FILE)
for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
    if ws.title != "read_me":
        for cell in ws[1]:
            if cell.value is not None:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F3A22")
    for col in ws.columns:
        w = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(w + 2, 12), 46)
wb.save(OUT_FILE)

print(f"Saved -> {OUT_FILE}")
print("\nFragility ranking:")
print(index_tab[["Crop", "Fragility score (0-100)", "Risk band"]].to_string(index=False))