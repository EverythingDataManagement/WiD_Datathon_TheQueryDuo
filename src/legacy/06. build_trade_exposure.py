#!/usr/bin/env python3
# =====================================================================
#  build_trade_exposure.py  —  Query Queens, GROW track
#  Turns the resilience finding into named, exposed importers.
#  For each FRAGILE crop, takes its #1 producer and finds:
#    - who imports the most FROM that producer (the exposed countries)
#    - how DEPENDENT each importer is (share of its imports of that crop
#      that comes from the one fragile supplier)
#  A crop with no backfill (resilience) + a highly dependent importer
#  = a specifically exposed country. That is the payoff slide.
#
#  Reads:  Trade_DetailedTradeMatrix_E_All_Data (updated).xlsx  (required)
#          Master_clean.xlsx  (optional; to auto-pick each crop's #1 producer)
#  Writes: Trade_Exposure.xlsx
#
#  Colab setup (run first):  !pip install openpyxl --quiet
# =====================================================================

import os
import pandas as pd
import numpy as np

# ---------- SETTINGS ----------
TRADE_FILE  = "Trade_DetailedTradeMatrix_E_All_Data (updated).xlsx"
TRADE_SHEET = "Sheet 1 - Trade_DetailedTradeMa"   # resolved tolerantly below
MASTER_FILE = "Master_clean.xlsx"                 # optional, to derive #1 producers
OUT_FILE    = "Trade_Exposure.xlsx"

TRADE_YEAR  = 2022    # a recent year with good trade coverage (change if needed)
TOP_N_IMPORTERS = 10  # how many exposed importers to list per crop

# Crop -> the single primary traded item name in the trade file (Option A: one form per crop).
# Rice is measured as milled rice (the dominant traded form). Sugarcane is not traded as cane.
CROP_TRADE_ITEM = {
    "Rice": "Rice, milled",
    "Maize": "Maize (corn)",
    "Wheat": "Wheat",
    "Soybeans": "Soya beans",
    "Potatoes": "Potatoes",
    "Cassava": "Cassava, fresh",
    "Sorghum": "Sorghum",
    "Millet": "Millet",
    "Barley": "Barley",
}
# The fragile crops from the resilience analysis (ratio < ~1). Focus the exposure story here.
FRAGILE_CROPS = ["Rice", "Maize", "Wheat", "Soybeans", "Cassava", "Potatoes"]

# Fallback #1 producers (used only if Master_clean.xlsx is not present).
# These match the resilience output. Names must match the TRADE file's country names.
FALLBACK_TOP_PRODUCER = {
    "Rice": "India", "Maize": "United States of America", "Wheat": "China, mainland",
    "Soybeans": "Brazil", "Cassava": "Nigeria", "Potatoes": "China, mainland",
    "Sorghum": "United States of America", "Millet": "India", "Barley": "Russian Federation",
}

KEY_TRADE_COLS = ["Reporter Countries", "Partner Countries", "Item", "Element"]


# =====================================================================
# LOADER — resolves the sheet, auto-detects the header row, validates the year
# =====================================================================
def load_trade(path, sheet, year):
    """Load only the needed trade columns, but fail with a message that says what to fix.
    Auto-detects whether the real header is on row 0 or row 1 (the 'reviewed/updated'
    FAOSTAT exports carry a title row above the header), resolves the sheet name
    tolerantly, and checks that the requested year column exists."""
    if not os.path.exists(path):
        here = sorted(f for f in os.listdir(".") if f.lower().endswith((".xlsx", ".xls")))
        lines = [f"Could not find '{path}' in the working directory:", f"    {os.getcwd()}"]
        if here:
            lines += ["Excel files that ARE here right now:", "    " + ", ".join(here),
                      f"Fix: set TRADE_FILE to one of those, or re-upload '{path}'."]
        else:
            lines += ["There are no Excel files in this directory at all.",
                      "In Colab, uploaded files vanish when the runtime restarts — re-upload and run again."]
        raise FileNotFoundError("\n".join(lines))

    xls = pd.ExcelFile(path)
    # resolve sheet: exact (case-insensitive) -> prefix -> first sheet
    match = next((s for s in xls.sheet_names if s.lower() == sheet.lower()), None)
    if match is None:
        match = next((s for s in xls.sheet_names if s.lower().startswith(sheet.lower()[:18])), None)
    if match is None:
        match = xls.sheet_names[0]

    # detect header row by which of the first two rows exposes the key columns
    header_row = None
    for h in (0, 1):
        cols = pd.read_excel(path, sheet_name=match, header=h, nrows=0).columns
        if all(k in cols for k in KEY_TRADE_COLS):
            header_row = h
            break
    if header_row is None:
        seen = list(pd.read_excel(path, sheet_name=match, header=0, nrows=0).columns)[:12]
        raise ValueError(
            f"Could not find the expected trade columns {KEY_TRADE_COLS} in the first two rows "
            f"of sheet '{match}'.\nColumns seen (header=0): {seen}\n"
            f"Fix: check TRADE_SHEET / the file is the Detailed Trade Matrix export.")

    all_cols = pd.read_excel(path, sheet_name=match, header=header_row, nrows=0).columns
    ycol = f"Y{year}"
    if ycol not in all_cols:
        yrs = [c for c in all_cols if c.startswith("Y") and c[1:].isdigit()]
        span = f"{yrs[0]}..{yrs[-1]}" if yrs else "none found"
        raise ValueError(f"Year column '{ycol}' is not in the trade file (available: {span}).\n"
                         f"Fix: set TRADE_YEAR to a year that exists.")

    df = pd.read_excel(path, sheet_name=match, header=header_row, usecols=KEY_TRADE_COLS + [ycol])

    # TRUNCATION CHECK: the full Detailed Trade Matrix has ~190 reporters and millions of
    # rows. Excel caps a sheet at 1,048,576 rows, so an .xlsx export is cut off mid-alphabet.
    # A round-number row count or a tiny reporter set means the file is incomplete.
    n_rows, n_reporters = len(df), df["Reporter Countries"].nunique()
    if n_rows >= 999_999 or n_reporters < 100:
        reps = sorted(df["Reporter Countries"].dropna().unique())
        print("\n*** WARNING: the trade file looks TRUNCATED — results will be incomplete ***")
        print(f"    rows read: {n_rows:,} | distinct reporter countries: {n_reporters}")
        if reps:
            print(f"    reporters run {reps[0]!r} .. {reps[-1]!r} (should span the whole alphabet)")
        print("    Fix: use the FAOSTAT 'Normalized' CSV bulk download (no row limit), or in the")
        print("    FAOSTAT interface pre-filter to the needed items + 'Import quantity' before export.\n")

    return df, match, header_row, ycol


# =====================================================================
# 1. Figure out each fragile crop's #1 PRODUCER (from the master if available)
# =====================================================================
top_producer = dict(FALLBACK_TOP_PRODUCER)
if os.path.exists(MASTER_FILE):
    try:
        mx = pd.ExcelFile(MASTER_FILE)
        msheet = next((s for s in mx.sheet_names if s.lower() == "master"), mx.sheet_names[0])
        m = pd.read_excel(MASTER_FILE, sheet_name=msheet)
        latest = m["Year"].max()
        for crop in FRAGILE_CROPS:
            s = m[(m["Crop"] == crop) & (m["Year"] == latest)].groupby("Country")["Production_tonnes"].sum()
            if len(s):
                top_producer[crop] = s.idxmax()
    except Exception as e:
        print(f"Note: could not derive top producers from {MASTER_FILE} ({e}); using the fallback list.")
print("Top producers used:", {c: top_producer[c] for c in FRAGILE_CROPS})

# =====================================================================
# 2. Load ONLY the trade rows we need (import quantity, our items) to keep it light
# =====================================================================
tr, trade_sheet, trade_header, ycol = load_trade(TRADE_FILE, TRADE_SHEET, TRADE_YEAR)
tr = tr.rename(columns={"Reporter Countries": "Importer", "Partner Countries": "Exporter",
                        ycol: "qty_t"})
# In the Detailed Trade Matrix, "Import quantity" reported by the Importer, Partner = the source.
# Case-insensitive match: FAOSTAT versions vary ("Import quantity" vs "Import Quantity").
tr = tr[tr["Element"].astype(str).str.strip().str.lower() == "import quantity"].copy()
tr["qty_t"] = pd.to_numeric(tr["qty_t"], errors="coerce")
tr = tr.dropna(subset=["qty_t"])
tr = tr[tr["qty_t"] > 0]

# drop aggregate reporters/partners (FAOSTAT rolls some up)
AGG = ["World","Africa","Americas","Asia","Europe","Oceania","European Union (27)",
       "Least Developed Countries","Net Food Importing Developing Countries",
       "Small Island Developing States","Land Locked Developing Countries",
       "Low Income Food Deficit Countries"]
tr = tr[~tr["Importer"].isin(AGG) & ~tr["Exporter"].isin(AGG)]

# =====================================================================
# 3. For each fragile crop: who imports MOST from its #1 producer, and how dependent
# =====================================================================
exposure_rows = []
skipped = []
for crop in FRAGILE_CROPS:
    item = CROP_TRADE_ITEM.get(crop)
    producer = top_producer.get(crop)
    if item is None or producer is None:
        skipped.append((crop, "no item/producer mapping"))
        continue
    ci = tr[tr["Item"] == item]
    if ci.empty:
        skipped.append((crop, f"no '{item}' import rows for {TRADE_YEAR} (check the trade item name)"))
        continue
    # each importer's TOTAL imports of this crop (from anywhere) -> for the dependency %
    tot_by_importer = ci.groupby("Importer")["qty_t"].sum()
    # imports specifically FROM the fragile top producer
    from_top = ci[ci["Exporter"] == producer].groupby("Importer")["qty_t"].sum()
    df = pd.DataFrame({"from_top_producer_t": from_top}).dropna()
    if df.empty:
        skipped.append((crop, f"no imports recorded from '{producer}' (check the producer name)"))
        continue
    df["total_imports_t"] = tot_by_importer.reindex(df.index)
    df["dependency_pct"] = (df["from_top_producer_t"] / df["total_imports_t"] * 100).round(0)
    df = df.sort_values("from_top_producer_t", ascending=False).head(TOP_N_IMPORTERS)
    for importer, row in df.iterrows():
        exposure_rows.append({
            "Crop": crop, "Top producer (at risk)": producer, "Traded as": item,
            "Exposed importer": importer,
            "Imports from producer (kt)": round(row["from_top_producer_t"] / 1000, 1),
            "Their total imports (kt)": round(row["total_imports_t"] / 1000, 1),
            "Dependency on this producer %": row["dependency_pct"],
        })

# build with explicit columns so an empty result still has a schema (no KeyError below)
EXP_COLS = ["Crop", "Top producer (at risk)", "Traded as", "Exposed importer",
            "Imports from producer (kt)", "Their total imports (kt)", "Dependency on this producer %"]
exposure = pd.DataFrame(exposure_rows, columns=EXP_COLS)

if skipped:
    print("\nWARNING - crops with no exposure rows (name mismatch vs the trade file, or no trade):")
    for c, why in skipped:
        print(f"   {c}: {why}")

# =====================================================================
# 4. A tight "most exposed" shortlist: high volume AND high dependency
# =====================================================================
short = exposure[exposure["Dependency on this producer %"] >= 40].copy()
short = short.sort_values(["Crop", "Dependency on this producer %"], ascending=[True, False])

# =====================================================================
# 5. WRITE
# =====================================================================
readme = pd.DataFrame([
    ["Query Queens - TRADE EXPOSURE", ""],
    ["Turns the resilience finding into NAMED exposed importers.", ""],
    ["", ""],
    ["For each fragile crop, we take its #1 producer and find which countries import most", ""],
    ["FROM that producer, and how dependent each is (share of their imports of that crop", ""],
    ["that comes from the one supplier). A crop with no backfill (resilience) + a highly", ""],
    ["dependent importer = a specifically exposed country.", ""],
    ["", ""],
    [f"Trade year: {TRADE_YEAR}. Rice measured as milled rice (dominant traded form).", ""],
    ["Import quantity, tonnes. Sugarcane excluded (traded as sugar, not cane).", ""],
    [f"Loaded sheet '{trade_sheet}' (header row {trade_header}).", ""],
    ["", ""],
    ["Tabs", ""],
    ["exposure", "top importers from each fragile crop's #1 producer, with dependency %"],
    ["most_exposed", "shortlist: importers >=40% dependent on a single fragile supplier"],
    ["", ""],
    ["Limitation", "This is who trades WITH the producer today; it does not model whether"],
    ["", "they could re-source (the resilience tab already shows the world can't backfill)."],
], columns=["A", "B"])

with pd.ExcelWriter(OUT_FILE, engine="openpyxl") as xl:
    readme.to_excel(xl, sheet_name="read_me", index=False, header=False)
    exposure.to_excel(xl, sheet_name="exposure", index=False)
    short.to_excel(xl, sheet_name="most_exposed", index=False)

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
wb = load_workbook(OUT_FILE)
for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
    if ws.title != "read_me":
        for cell in ws[1]:
            if cell.value is not None:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F3A22")
    for col in ws.columns:
        w = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(w + 2, 12), 45)
wb.save(OUT_FILE)

print(f"\nSaved -> {OUT_FILE}")
print(f"exposure rows: {len(exposure)} | most-exposed shortlist: {len(short)}")