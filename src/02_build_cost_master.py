#!/usr/bin/env python3
# =====================================================================
#  build_cost_master.py  —  Query Queens, GROW track
#  Joins FAOSTAT Land Use (RL) + Fertilizers by Nutrient (RFN) into ONE
#  country-year cost master. Grain = one row per Country x Year, with:
#     - cropland & forest area (1000 ha)     + their source flags
#     - N + P + K fertilizer intensity (kg per ha of cropland) + flag
#  These sit side by side so the land-cost and yield-cost arms read from
#  a single file. Both sources are national (not crop-level), which is
#  why the grain is country x year, not country x crop x year.
#
#  UPDATE (this version): the pipeline now KEEPS Item, Element and the
#  per-year source FLAG. Because the joined cost_master is pivoted
#  (Item collapses into column names, Element is uniform per source),
#  Item/Element/Flag are exposed two ways:
#     - cost_master  : gains cropland_flag, forest_flag, fertilizer_flag
#     - source_detail: NEW tidy sheet, one row per Country x Year x Item
#                      showing Item, Element, Unit, Value, Flag directly
#     - flag_summary : NEW, the flag mix per Item at a glance
#
#  FAOSTAT year columns come in triples: Y#### (value), Y####F (flag),
#  Y####N (note). We read value + flag; the note column is ignored.
#
#  Colab setup (run first in a cell):  !pip install openpyxl --quiet
# =====================================================================

import os
import pandas as pd
import numpy as np

# ---------- SETTINGS ----------
LAND_FILE  = "Inputs_LandUse_E_All_Data (Reviewed).xlsx"              # RL
FERT_FILE  = "Inputs_FertilizersNutrient_E_All_Data (Reviewed).xlsx"  # RFN
# These reviewed files have a TITLE row above the real headers, so the header is row 2
# (header=1, zero-indexed). Sheet names below match the reviewed files.
LAND_SHEET = "Sheet 1 - Inputs_LandUse_E_All_"
FERT_SHEET = "Sheet 1 - Inputs_FertilizersNut"
HEADER_ROW = 1
OUT_FILE  = "Cost_Master.xlsx"
START_YEAR, END_YEAR = 2000, 2024

# RL codes
AREA_ELEMENT = 5110          # "Area" element, in 1000 ha
CROPLAND, FOREST = 6620, 6646

# Exact FAOSTAT aggregate / non-country names to drop (never keyword-match)
AGGREGATES = [
    "World","Africa","Americas","Asia","Europe","Oceania","Eastern Africa","Middle Africa",
    "Northern Africa","Southern Africa","Western Africa","Northern America","Central America",
    "Caribbean","South America","Latin America and the Caribbean","Sub-Saharan Africa",
    "Central Asia","Eastern Asia","Southern Asia","South-eastern Asia","Western Asia",
    "Eastern Europe","Northern Europe","Southern Europe","Western Europe",
    "Australia and New Zealand","Melanesia","Micronesia","Polynesia","European Union (27)",
    "Least Developed Countries (LDCs)","Land Locked Developing Countries (LLDCs)",
    "Low Income Food Deficit Countries (LIFDCs)","Net Food Importing Developing Countries (NFIDCs)",
    "Small Island Developing States (SIDS)","High-income economies","Low-income economies",
    "Lower-middle-income economies","Upper-middle-income economies","Middle-income economies",
]

log = []
def rec(step, what, count):
    log.append({"Step": step, "What happened": what, "Count": count})
    print(f"[{step}] {what}: {count}")


def melt_with_flags(df, id_cols, start, end):
    """FAOSTAT wide -> long, carrying BOTH the per-year value and its flag.

    Each year is three columns: Y#### (value), Y####F (flag), Y####N (note).
    We melt value cols and flag cols separately, then re-join them on a
    stable source-row id so value<->flag alignment is exact even if two
    source rows happen to be identical. The note column is dropped.
    """
    df = df.reset_index(drop=True).copy()
    df["_rid"] = df.index
    val_cols  = [c for c in df.columns
                 if c.startswith("Y") and c[1:].isdigit()]                       # Y2000
    flag_cols = [c for c in df.columns
                 if c.startswith("Y") and c.endswith("F") and c[1:-1].isdigit()]  # Y2000F

    v = df.melt(id_vars=id_cols + ["_rid"], value_vars=val_cols,
                var_name="Yv", value_name="value")
    v["Year"] = v["Yv"].str[1:].astype(int)

    f = df.melt(id_vars=["_rid"], value_vars=flag_cols,
                var_name="Yf", value_name="flag")
    f["Year"] = f["Yf"].str[1:-1].astype(int)

    out = (v.drop(columns="Yv")
             .merge(f.drop(columns="Yf"), on=["_rid", "Year"], how="left")
             .drop(columns="_rid"))
    return out[out["Year"].between(start, end)].copy()


def combine_flags(s):
    """Fertilizer intensity is a SUM of N+P+K, so its 'flag' is the set of the
    three nutrient flags. Return distinct flags joined by ';' (e.g. 'A' or 'A;E'),
    so a mixed official/imputed sum is visible rather than hidden."""
    vals = sorted({str(x) for x in s.dropna() if str(x).strip()})
    return ";".join(vals) if vals else np.nan


# =====================================================================
# 1. LAND USE (RL) -> tidy country-year cropland & forest (1000 ha) + flags
# =====================================================================
if not os.path.exists(LAND_FILE):
    raise FileNotFoundError(f"{LAND_FILE} not found - upload it and check the name.")
rl = pd.read_excel(LAND_FILE, sheet_name=LAND_SHEET, header=HEADER_ROW)
rec("1. Land", "Raw RL rows", len(rl))
rl = rl[(rl["Element Code"] == AREA_ELEMENT) & (rl["Item Code"].isin([CROPLAND, FOREST]))]

# long form keeps Item / Element / Unit / value / flag
land_long = melt_with_flags(
    rl, ["Area", "Item Code", "Item", "Element", "Unit"], START_YEAR, END_YEAR)

# wide values (as before)
land_val = land_long.pivot_table(index=["Area", "Year"], columns="Item Code",
                                 values="value", aggfunc="first").reset_index()
land_val = land_val.rename(columns={CROPLAND: "cropland_1000ha", FOREST: "forest_1000ha"})

# wide flags (new) - one flag per item, straight through
land_flag = land_long.pivot_table(index=["Area", "Year"], columns="Item Code",
                                  values="flag", aggfunc="first").reset_index()
land_flag = land_flag.rename(columns={CROPLAND: "cropland_flag", FOREST: "forest_flag"})

land = land_val.merge(land_flag, on=["Area", "Year"], how="left")
rec("1. Land", "Land country-year rows", len(land))


# =====================================================================
# 2. FERTILIZERS (RFN) -> tidy country-year N+P+K intensity (kg/ha) + flag
# =====================================================================
if not os.path.exists(FERT_FILE):
    raise FileNotFoundError(f"{FERT_FILE} not found - upload it and check the name.")
rf = pd.read_excel(FERT_FILE, sheet_name=FERT_SHEET, header=HEADER_ROW)
rec("2. Fertilizer", "Raw RFN rows", len(rf))
# match by name (robust across FAOSTAT versions): the 3 nutrients + per-cropland-area element
rf = rf[rf["Item"].str.contains("nitrogen|phosphate|potash", case=False, na=False) &
        rf["Element"].str.contains("per area of cropland", case=False, na=False)]

# long form keeps each nutrient as its own row, with its Item / Element / Unit / value / flag
fert_long = melt_with_flags(
    rf, ["Area", "Item Code", "Item", "Element", "Unit"], START_YEAR, END_YEAR)

# sum the three nutrients -> total NPK intensity per country-year
fert = (fert_long.groupby(["Area", "Year"])["value"].sum(min_count=1)
        .reset_index().rename(columns={"value": "fertilizer_kg_per_ha"}))
# combined flag across the three nutrients (distinct set)
fert_flag = (fert_long.groupby(["Area", "Year"])["flag"].apply(combine_flags)
             .reset_index().rename(columns={"flag": "fertilizer_flag"}))
fert = fert.merge(fert_flag, on=["Area", "Year"], how="left")
rec("2. Fertilizer", "Fertilizer country-year rows", len(fert))


# =====================================================================
# 3. JOIN on Country + Year (outer, so a country present in one file
#    but not the other is kept with blanks rather than dropped)
# =====================================================================
cost = land.merge(fert, on=["Area", "Year"], how="outer")
cost = cost.rename(columns={"Area": "Country"})

# drop FAOSTAT aggregates / non-country rows
before = cost["Country"].nunique()
cost = cost[~cost["Country"].isin(AGGREGATES)]
rec("3. Join", "Aggregates removed", before - cost["Country"].nunique())
rec("3. Join", "Countries in cost master", cost["Country"].nunique())

# order + sort (value then its flag, so each measure reads next to its provenance)
cost = cost[["Country", "Year",
             "cropland_1000ha", "cropland_flag",
             "forest_1000ha", "forest_flag",
             "fertilizer_kg_per_ha", "fertilizer_flag"]]
cost = cost.sort_values(["Country", "Year"]).reset_index(drop=True)
rec("3. Join", "FINAL cost-master rows", len(cost))


# =====================================================================
# 3b. SOURCE DETAIL (NEW) -> tidy long view so Item / Element / Flag
#     are visible per underlying figure (before cropland/forest/NPK
#     get collapsed into cost_master's columns).
# =====================================================================
land_sd = land_long.rename(columns={"Area": "Country"}).assign(Source="Land Use (RL)")
fert_sd = fert_long.rename(columns={"Area": "Country"}).assign(Source="Fertilizers by Nutrient (RFN)")
source_detail = pd.concat([land_sd, fert_sd], ignore_index=True)
source_detail = source_detail[~source_detail["Country"].isin(AGGREGATES)]
source_detail = source_detail.dropna(subset=["value"])          # keep rows that have a real figure
source_detail = source_detail.rename(columns={"value": "Value", "flag": "Flag"})
source_detail = source_detail[["Country", "Year", "Source",
                               "Item", "Element", "Unit", "Value", "Flag"]]
source_detail = source_detail.sort_values(
    ["Country", "Year", "Source", "Item"]).reset_index(drop=True)
rec("3b. Detail", "source_detail rows", len(source_detail))


# =====================================================================
# 4. COVERAGE / DATA-QUALITY per measure column
# =====================================================================
N = len(cost)
dq = []
for col in ["cropland_1000ha", "forest_1000ha", "fertilizer_kg_per_ha"]:
    s = cost[col]
    blank = int(s.isna().sum())
    zero = int((s == 0).sum())
    dq.append({"Column": col, "Blank count": blank, "Blank %": round(blank / N * 100, 2),
               "Zero count": zero, "Zero %": round(zero / N * 100, 2),
               "Completeness %": round(100 - blank / N * 100, 2)})
dq_df = pd.DataFrame(dq)


# =====================================================================
# 4b. FLAG SUMMARY (NEW) -> flag mix per Item, so imputed / estimated
#     shares are visible at a glance (see README for flag meanings).
# =====================================================================
flag_summary = (source_detail.assign(Flag=source_detail["Flag"].fillna("(none)"))
                .groupby(["Item", "Flag"]).size().reset_index(name="Row count"))
flag_summary = flag_summary.sort_values(
    ["Item", "Row count"], ascending=[True, False]).reset_index(drop=True)


# =====================================================================
# 5. READ_ME + REFERENCE
# =====================================================================
readme = pd.DataFrame([
    ["Query Queens - COST MASTER (land + fertilizer)", ""],
    ["Women in Data 2026 Datathon, GROW track", ""],
    ["", ""],
    ["What this file is", ""],
    ["One country-year table joining FAOSTAT Land Use (RL) and Fertilizers by Nutrient (RFN).", ""],
    ["Both sources are NATIONAL (not crop-level), so the grain is Country x Year.", ""],
    ["Rebuilt by build_cost_master.py. Do not hand-edit; fix the script and re-run.", ""],
    ["", ""],
    ["Sheets", ""],
    ["cost_master", "Joined Country x Year table; each measure now sits next to its source flag."],
    ["source_detail", "Tidy long view: one row per Country x Year x Item, with Item, Element, Unit, Value, Flag."],
    ["data_quality", "Blank / zero / completeness per measure column."],
    ["flag_summary", "Row counts per Item x Flag - the source-quality mix at a glance."],
    ["cleaning_log", "Step-by-step row counts from the build."],
    ["", ""],
    ["Column meanings (cost_master)", ""],
    ["cropland_1000ha", "Cropland area, 1000 hectares (RL item 6620)."],
    ["cropland_flag", "FAOSTAT source flag for the cropland figure (see legend below)."],
    ["forest_1000ha", "Forest land area, 1000 hectares (RL item 6646)."],
    ["forest_flag", "FAOSTAT source flag for the forest figure."],
    ["fertilizer_kg_per_ha", "N + P2O5 + K2O applied per hectare of cropland (RFN, summed)."],
    ["fertilizer_flag", "Distinct flags across the 3 nutrients, joined by ';' (e.g. 'A' or 'A;E')."],
    ["", ""],
    ["FAOSTAT flag legend (confirm against the source's own notes)", ""],
    ["A", "Official figure."],
    ["E", "Estimated value."],
    ["I", "Imputed value."],
    ["P", "Provisional value."],
    ["X", "Figure from an international organisation / source."],
    ["(none)", "No flag recorded for that cell."],
    ["", ""],
    ["How to use it", ""],
    ["Land-cost arm: compare cropland change vs forest change 2000->latest (cropland up + forest down = frontier clearing).", ""],
    ["Yield-cost arm: fertilizer_kg_per_ha is the input-intensity measure.", ""],
    ["Read the flag before trusting a latest-year value: I (imputed) / E (estimated) are weaker than A (official).", ""],
    ["Crop attribution stays INFERRED (from the production master's area/yield decomposition) - these files are national.", ""],
    ["", ""],
    ["Known limitations", ""],
    ["A. Latest FAOSTAT land/forest figures may be imputed - the flag columns now surface this; filter on them.", ""],
    ["B. Land data may run to a different latest year than fertilizer - the join keeps each year's real value.", ""],
    ["C. National forest nets planting against clearing (e.g. China afforestation), so it is not primary-forest loss.", ""],
    ["D. source_detail drops cells with no value; a flag with no value is not carried.", ""],
], columns=["A", "B"])


# =====================================================================
# 6. WRITE
# =====================================================================
with pd.ExcelWriter(OUT_FILE, engine="openpyxl") as xl:
    readme.to_excel(xl, sheet_name="read_me", index=False, header=False)
    cost.to_excel(xl, sheet_name="cost_master", index=False)
    source_detail.to_excel(xl, sheet_name="source_detail", index=False)
    dq_df.to_excel(xl, sheet_name="data_quality", index=False)
    flag_summary.to_excel(xl, sheet_name="flag_summary", index=False)
    pd.DataFrame(log).to_excel(xl, sheet_name="cleaning_log", index=False)

# light header styling
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
wb = load_workbook(OUT_FILE)
for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
    if ws.title != "read_me":
        for cell in ws[1]:
            if cell.value is not None:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F3A22")
    for col in ws.columns:
        w = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(w + 2, 12), 50)
wb.save(OUT_FILE)

print(f"\nSaved -> {OUT_FILE}")
print(f"cost master: {len(cost)} rows, {cost['Country'].nunique()} countries, {START_YEAR}-{END_YEAR}")
print(f"source detail: {len(source_detail)} rows (Country x Year x Item, with Item/Element/Flag)")